#!/usr/bin/env python
"""生成 2026-08-07 导师汇报 pptx（学术表述版）。
文案源：2026-08-07-advisor-report-draft.md。
运行：~/.venvs/slides/bin/python make-2026-08-07-advisor-pptx.py
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import os

# ---- 版面与配色 ----
SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.75)
CONTENT_W = SLIDE_W - 2 * MARGIN

BG_DARK = RGBColor(0x1C, 0x30, 0x4A)   # 深蓝底（标题页 / 强调块）
ACCENT  = RGBColor(0x2E, 0x6F, 0xC2)   # 主题蓝
CARD    = RGBColor(0xF3, 0xF5, 0xF9)   # 浅灰卡片
INK     = RGBColor(0x22, 0x2B, 0x36)   # 正文深色
MUTED   = RGBColor(0x7C, 0x86, 0x92)   # 次要灰
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
SOFT    = RGBColor(0xC9, 0xD6, 0xE8)   # 深底上的浅蓝灰
GREEN   = RGBColor(0x2F, 0x9E, 0x63)
AMBER   = RGBColor(0xC0, 0x84, 0x1F)
GRAY    = RGBColor(0x9A, 0xA3, 0xAD)

FONT_CJK = "Hiragino Sans GB"    # 冬青黑体：中文
FONT_LATIN = "Helvetica Neue"    # 拉丁字符

prs = Presentation()
prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
BLANK = prs.slide_layouts[6]


def style_run(run, size, bold=False, color=INK):
    f = run.font
    f.name = FONT_LATIN
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", FONT_CJK)


def add_box(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP):
    """paras: list of dict(text, size, bold, color, align, space_before, space_after, line)"""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = p.get("align", PP_ALIGN.LEFT)
        if p.get("space_before"):
            para.space_before = Pt(p["space_before"])
        if p.get("space_after"):
            para.space_after = Pt(p["space_after"])
        if p.get("line"):
            para.line_spacing = p["line"]
        run = para.add_run()
        run.text = p["text"]
        style_run(run, p["size"], p.get("bold", False), p.get("color", INK))
    return tb


def add_rect(slide, x, y, w, h, fill, rounded=True, radius=0.055):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, x, y, w, h)
    if rounded:
        shape.adjustments[0] = radius
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def new_slide():
    return prs.slides.add_slide(BLANK)


def set_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def content_header(slide, kicker, title, kicker_color=ACCENT):
    if kicker:
        add_box(slide, MARGIN, Inches(0.48), CONTENT_W, Inches(0.32),
                [dict(text=kicker, size=13, bold=True, color=kicker_color)])
    add_box(slide, MARGIN, Inches(0.82), CONTENT_W, Inches(0.55),
            [dict(text=title, size=25, bold=True, color=INK)])
    add_rect(slide, MARGIN, Inches(1.46), Inches(0.55), Inches(0.045), ACCENT, rounded=False)


def footer(slide, page_no):
    add_box(slide, SLIDE_W - MARGIN - Inches(0.6), Inches(7.06), Inches(0.6), Inches(0.3),
            [dict(text=str(page_no), size=9, color=MUTED, align=PP_ALIGN.RIGHT)])


def block_rows(slide, rows, y0, row_h, gap, label_w=Inches(1.4), body_size=14.5, width=None):
    """rows: list of (label, body, tag, emphasized)；tag 为 None 时不显示"""
    w = width if width is not None else CONTENT_W
    y = y0
    for label, body, tag, emph in rows:
        fill = BG_DARK if emph else CARD
        add_rect(slide, MARGIN, y, w, row_h, fill)
        lab_color = SOFT if emph else ACCENT
        body_color = WHITE if emph else INK
        tag_color = SOFT if emph else ACCENT
        pad = Inches(0.3)
        add_box(slide, MARGIN + pad, y, label_w, row_h,
                [dict(text=label, size=13, bold=True, color=lab_color)],
                anchor=MSO_ANCHOR.MIDDLE)
        paras = [dict(text=body, size=body_size, color=body_color, line=1.25)]
        if tag:
            paras.append(dict(text=tag, size=12, bold=True, color=tag_color,
                              space_before=6))
        add_box(slide, MARGIN + pad + label_w + Inches(0.15), y,
                w - 2 * pad - label_w - Inches(0.15), row_h,
                paras, anchor=MSO_ANCHOR.MIDDLE)
        y += row_h + gap


def side_figure(slide, filename):
    """研究点页右侧示意图；文件缺失时静默跳过"""
    path = os.path.join(os.path.dirname(os.path.abspath(__file__)), filename)
    if not os.path.exists(path):
        return
    from PIL import Image as _PILImage
    _w, _h = _PILImage.open(path).size
    fw = Inches(3.83)
    fh = int(fw * _h / _w)
    if fh > Inches(5.15):
        fh = Inches(5.15)
        fw = int(fh * _w / _h)
    slide.shapes.add_picture(path, int(Inches(8.75) + (Inches(3.83) - fw) / 2), Inches(1.75), fw, fh)


ROWS_W = Inches(7.75)  # 带示意图页的行宽

# ================= 第 1 页：标题 =================
s = new_slide()
add_rect(s, 0, 0, SLIDE_W, SLIDE_H, BG_DARK, rounded=False)
add_box(s, Inches(1.5), Inches(2.35), SLIDE_W - Inches(3), Inches(1.0),
        [dict(text="Omni-Anything", size=48, bold=True, color=WHITE, align=PP_ALIGN.CENTER)])
add_box(s, Inches(1.5), Inches(3.45), SLIDE_W - Inches(3), Inches(0.5),
        [dict(text="为智能体提供前台双工交互能力的推理服务系统", size=21, color=SOFT, align=PP_ALIGN.CENTER)])
add_rect(s, (SLIDE_W - Inches(0.7)) / 2, Inches(4.15), Inches(0.7), Inches(0.05), ACCENT, rounded=False)
add_box(s, Inches(2.2), Inches(4.5), SLIDE_W - Inches(4.4), Inches(0.9),
        [dict(text="支撑「前台双工交互 + 后台智能体协同」负载的高效推理服务",
              size=15, color=WHITE, align=PP_ALIGN.CENTER)])
set_notes(s, "智能体的推理与工具能力已经成熟，但与用户的交互仍以一问一答为主。本项目为智能体配备双工交互前台——前台模型持续接收输入、实时生成输出，后台智能体异步执行任务、结果实时写回——并研究支撑这一负载的推理服务系统。")

# ================= 第 2 页：项目总览（大图） =================
s = new_slide()
content_header(s, "项目总览", "为 agentic 系统构建实时双工交互底座")
add_box(s, MARGIN, Inches(1.58), CONTENT_W, Inches(0.62),
        [dict(text="JiuwenSwarm 等 agentic 系统已经能替人完成复杂任务，但触达用户仍以文字问答为主。本项目为它们构建实时语音双工交互的推理底座：这层系统决定双工交互能否规模化上线，也决定后台智能体的成果能否被及时说出口。",
              size=13, color=INK, line=1.25)])
FIG_OV = os.path.join(os.path.dirname(os.path.abspath(__file__)), "figure-overview-jiuwenswarm.png")
if os.path.exists(FIG_OV):
    from PIL import Image as _PILImage
    _w, _h = _PILImage.open(FIG_OV).size
    fig_h = Inches(4.72)
    fig_w = int(fig_h * _w / _h)
    if fig_w > CONTENT_W:
        fig_w = CONTENT_W
        fig_h = int(fig_w * _h / _w)
    s.shapes.add_picture(FIG_OV, int((SLIDE_W - fig_w) / 2), Inches(2.3), fig_w, fig_h)
footer(s, 2)
set_notes(s, "先讲价值再讲图：JiuwenSwarm 这样的系统已经很能干活，但和用户的接触方式还是文字问答。我们给它配上实时语音双工前台，而这件事能不能规模化上线、后台干完的活能不能及时说出口，全部取决于底下这层推理系统——这就是本项目的位置。然后看图：左边用户持续说话随时打断，中间是我们的运行时，右边 JiuwenSwarm 在后台组队干活、结果异步写回；下半的时间轴是负载本质——前台刚性周期，后台弹性任务。三个标注是三个研究点的由来，每一条都写了「为什么这是个问题」。带着这三个问题往下看。")

# ================= 第 3 页：研究背景 =================
s = new_slide()
content_header(s, "研究背景", "智能体双工交互已进入产品，系统层研究刚刚起步")
block_rows(s, [
    ("交互范式演进",
     "智能体的推理与工具调用能力已趋成熟，但其与用户的交互仍以「请求—响应」为主。新一代产品为智能体配备了双工交互前台：前台模型持续接收输入、同时实时生成输出（如实时语音），后台智能体异步执行任务并将结果写回对话。OpenAI 与字节跳动今年相继发布此类产品。",
     None, False),
    ("研究现状与空白",
     "上述产品的服务侧方案均未公开；学术界对该负载的系统研究刚刚起步，「前台双工交互 + 后台结果写回」的服务问题仍缺乏公开的系统性方案。",
     None, False),
    ("研究目标",
     "为智能体的前台双工交互构建推理服务系统：在实时性约束下，最大化单位 GPU 资源的并发会话承载能力。",
     None, True),
], y0=Inches(1.85), row_h=Inches(1.42), gap=Inches(0.24))
footer(s, 3)
set_notes(s, "智能体的能力已经成熟，交互形态正在升级：新一代产品给智能体配上了双工交互前台，前台持续听说，后台智能体干活、结果实时写回，OpenAI 和字节的产品已经验证了这一方向。但服务侧方案都不公开，学术研究刚刚起步。本项目的研究目标是：在实时性约束下，最大化单位 GPU 资源可承载的并发会话数。")

# ================= 第 4 页：任务形态 =================
s = new_slide()
content_header(s, "任务形态", "双工模型的任务形态：周期性连续请求，每周期须按时完成")
add_box(s, MARGIN, Inches(1.60), CONTENT_W, Inches(0.58),
        [dict(text="双工模型不是「一次请求、一次回答」：会话期间模型以固定周期持续运转，每个周期都要在 deadline 前完成一轮计算；后台智能体的任务调用与结果写回穿插其中。",
              size=13, color=INK, line=1.2)])
FIG = os.path.join(os.path.dirname(os.path.abspath(__file__)), "figure-duplex-workload-cn.png")
if os.path.exists(FIG):
    from PIL import Image as _PILImage
    _w, _h = _PILImage.open(FIG).size
    fig_h = Inches(5.05)
    fig_w = int(fig_h * _w / _h)
    if fig_w > CONTENT_W:
        fig_w = CONTENT_W
        fig_h = int(fig_w * _h / _w)
    s.shapes.add_picture(FIG, int((SLIDE_W - fig_w) / 2), Inches(2.28), fig_w, fig_h)
footer(s, 4)
set_notes(s, "先建立任务形态的直觉：普通聊天或智能体请求是回合制的，回合之间有空闲，系统可以利用空闲腾挪状态。双工会话完全不同——它是周期性的连续请求，每个周期都有硬性时限，会话全程没有空闲，状态必须常驻且持续增长。再叠加一层：前台在跟用户实时交互的同时，还会调用后台智能体或工具，结果要写回对话。后面的三个研究点，都建立在这个任务形态之上。")

# ================= 第 5 页：核心挑战 =================
s = new_slide()
content_header(s, "问题与挑战", "现有推理系统与持续交互负载的根本失配")
block_rows(s, [
    ("挑战 1｜状态容量",
     "会话状态常驻且持续增长：上下文状态（KV 缓存）随会话时长单调增长，且每个交互周期都必须访问、无法换出。实测表明，GPU 计算远未饱和时，会话已因显存耗尽而失败。",
     "→ 对应研究点 1", False),
    ("挑战 2｜执行效率",
     "周期计算粒度小、频率高：每个交互周期的有效计算量很小，而通用引擎将其视为动态负载，调度、内核启动与内存规划等固定开销按周期重复支付，占比显著。",
     "→ 对应研究点 2", False),
    ("挑战 3｜链路协同",
     "端到端实时性由多个环节共同决定：一次双工交互跨越网络传输、前台模型推理、后台智能体调用与客户端播放，各环节由独立子系统按局部目标优化、彼此状态不可见，实时时限只能靠各环节独立预留余量保障。",
     "→ 对应研究点 3", False),
], y0=Inches(1.85), row_h=Inches(1.5), gap=Inches(0.2), label_w=Inches(1.75))
footer(s, 5)
set_notes(s, "这类负载与现有系统存在三项根本失配：其一，会话状态常驻且持续增长，显存容量先于计算能力成为瓶颈；其二，每周期计算量小而动态执行开销大；其三，端到端实时性横跨传输、推理、后台调用与播放多个环节，而各环节孤立优化、互不知情。三项挑战分别由三个研究点针对性解决——和总览大图上的三个标注一一对应。")

# ================= 第 6 页：研究内容总览 =================
s = new_slide()
content_header(s, "研究内容", "三个研究点")
add_rect(s, MARGIN, Inches(1.8), CONTENT_W, Inches(0.62), BG_DARK)
add_box(s, MARGIN, Inches(1.8), CONTENT_W, Inches(0.62),
        [dict(text="总体目标：在实时性约束下，最大化单位 GPU 资源的并发会话承载能力", size=15, bold=True,
              color=WHITE, align=PP_ALIGN.CENTER)], anchor=MSO_ANCHOR.MIDDLE)

cards = [
    ("研究点 1", "会话状态容量扩展",
     "利用交互周期的可预测性，以闲置传输带宽换取等效显存容量，无损保全上下文", "● 进行中，有实验支撑", GREEN),
    ("研究点 2", "周期执行静态特化",
     "利用周期任务输入形状固定的特性，将逐周期动态调度转化为静态执行计划的直接重放", "● 进行中，有实验支撑", GREEN),
    ("研究点 3", "端到端链路联合调度",
     "将网络传输、前台推理、后台智能体调用与播放纳入同一控制环，共同守护端到端实时性", "○ 立项设想", GRAY),
]
gap = Inches(0.3)
card_w = (CONTENT_W - 2 * gap) / 3
card_y, card_h = Inches(2.72), Inches(3.75)
for i, (no, name, desc, status, scolor) in enumerate(cards):
    x = MARGIN + i * (card_w + gap)
    add_rect(s, x, card_y, card_w, card_h, CARD)
    pad = Inches(0.3)
    add_box(s, x + pad, card_y + Inches(0.35), card_w - 2 * pad, Inches(2.3),
            [dict(text=no, size=13, bold=True, color=ACCENT, space_after=6),
             dict(text=name, size=20, bold=True, color=INK, space_after=12),
             dict(text=desc, size=14, color=INK, line=1.3)])
    add_box(s, x + pad, card_y + card_h - Inches(0.65), card_w - 2 * pad, Inches(0.4),
            [dict(text=status, size=13, bold=True, color=scolor)])
footer(s, 6)
set_notes(s, "项目的总体目标是在实时性约束下，用同样的 GPU 资源承载更多并发会话。研究点一扩展状态容量、研究点二消除动态执行开销，两者都已进入实验验证并各有一页实验结果；研究点三把端到端链路纳入联合调度，处于立项设想。")

# ================= 第 7 页：研究点 1 =================
s = new_slide()
content_header(s, "研究点 1｜进行中", "会话状态虚拟化：以闲置传输带宽无损扩展等效显存容量", kicker_color=GREEN)
block_rows(s, [
    ("研究问题",
     "并发会话的上下文状态持续增长，显存容量先于计算能力成为瓶颈；会话因显存耗尽而失败时，GPU 利用率仍处于低位。",
     None, False),
    ("关键洞察",
     "交互负载具有严格的周期性：每路会话下一周期的计算时刻与所需状态可精确预知，这使状态迁移可以按时间表调度。",
     None, False),
    ("技术路线",
     "将近期不活跃的上下文状态存放于主机内存，按交互周期提前预取回显存、用后即时释放；迁移经独立传输通道进行，与计算完全重叠，不占用实时时限。",
     None, False),
    ("预期与进展",
     "瓶颈与方案收益均已由真机实验验证（下页）；上下文无损——现有方案依赖有损截断。",
     None, True),
], y0=Inches(1.75), row_h=Inches(1.18), gap=Inches(0.16), width=ROWS_W)
side_figure(s, "figure-rp1-idea.png")
footer(s, 7)
set_notes(s, "核心观察是交互负载的严格周期性，使状态迁移可以按时间表调度：不活跃的上下文状态放在主机内存，按周期预取、用后释放，迁移与计算重叠。瓶颈和方案收益都已经在真机上验证，下一页看实验结果；关键差异是上下文无损，现有方案靠截断历史换容量。")

# ================= 第 8 页：研究点 1 · 实验验证 =================
s = new_slide()
content_header(s, "研究点 1｜实验验证", "同负载公平对照：驻留行为的差异与来源", kicker_color=GREEN)
IMG_BL = os.path.join(os.path.dirname(os.path.abspath(__file__)), "rp1-baseline-kv-crop.png")
IMG_CV = os.path.join(os.path.dirname(os.path.abspath(__file__)), "rp1-conveyer-kv-crop.png")
if os.path.exists(IMG_BL) and os.path.exists(IMG_CV):
    from PIL import Image as _PILImage
    add_box(s, MARGIN, Inches(1.78), Inches(7.3), Inches(0.3),
            [dict(text="基线：KV 全程常驻，随会话时长单调增长（8 路会话 · 白盒时间线实测）", size=12.5, bold=True, color=INK)])
    _w, _h = _PILImage.open(IMG_BL).size
    s.shapes.add_picture(IMG_BL, MARGIN, Inches(2.12), Inches(7.3), int(Inches(7.3) * _h / _w))
    y2 = Inches(2.12) + int(Inches(7.3) * _h / _w) + Inches(0.28)
    add_box(s, MARGIN, y2, Inches(7.3), Inches(0.3),
            [dict(text="本方案：仅计算窗口驻留、用后即时释放——以带宽换显存，瞬时占用最小化", size=12.5, bold=True, color=INK)])
    _w2, _h2 = _PILImage.open(IMG_CV).size
    s.shapes.add_picture(IMG_CV, MARGIN, y2 + Inches(0.34), Inches(7.3), int(Inches(7.3) * _h2 / _w2))
px, pw = Inches(8.4), Inches(4.18)
add_rect(s, px, Inches(1.78), pw, Inches(5.08), CARD)
pad = Inches(0.32)
add_box(s, px + pad, Inches(2.05), pw - 2 * pad, Inches(4.6),
        [dict(text="收益来源", size=13, bold=True, color=GREEN),
         dict(text="会话状态全量仍随时长增长，但常驻显存的部分被压缩为正在计算的工作集：交互周期严格可预知，状态按时间表预取回显存、用后立即释放，增长的尾部驻留主机内存，搬运与计算完全重叠。", size=12.5, color=INK, line=1.3, space_before=6),
         dict(text="量级（初步实验）", size=13, bold=True, color=GREEN, space_before=14),
         dict(text="稳态驻留约为基线的三分之一；交付质量零恶化，上下文无损。", size=12.5, color=INK, line=1.3, space_before=6),
         dict(text="实验设置与边界", size=13, bold=True, color=GREEN, space_before=14),
         dict(text="vLLM 0.23 + Qwen2.5-Omni-7B，RTX 3090，2 秒周期，8 路会话；消费级 GPU 初步实验，验证机制与洞察，生产容量待正式扫描。", size=12, color=MUTED, line=1.3, space_before=6)])
footer(s, 8)
set_notes(s, "这一页用真实的白盒时间线说话。上图是基线：八路会话的 KV 驻留全程常驻、一路单调上涨——会话活得越久占得越多，显存注定耗尽。下图是我们的方案：每路会话只在自己的计算窗口驻留，用完立刻释放，锯齿有界，跑多久都不再涨。收益的来源写在右边：不是什么魔法，是交互周期严格可预知，状态搬运可以排进时间表、藏进计算。量级上稳态驻留约为基线三分之一——初步实验看量级；交付零恶化、上下文无损。")

# ================= 第 9 页：研究点 2 =================
s = new_slide()
content_header(s, "研究点 2｜进行中", "面向周期性交互负载的静态执行特化", kicker_color=AMBER)
block_rows(s, [
    ("研究问题",
     "交互周期内的有效计算量小、执行形状高度重复；通用引擎将每个周期视为全新任务，调度、内核启动与内存规划等固定开销按周期重复支付，占比显著。",
     None, False),
    ("关键洞察",
     "双工负载的周期任务输入是固定形状：每周期输入长度由帧率与周期时长决定、恒定不变，批次组成亦可预知，执行形状可归约为有限模板族——而通用引擎按「有无新输入」分类执行路径，结构性看不见这类稳定性。",
     None, False),
    ("技术路线",
     "对每类模板预先构建静态执行计划（整步计算图捕获、注意力元数据固化、工作区预分配），运行时直接重放，以消除重复的动态决策开销；形状失配时自动回退通用路径。",
     None, False),
    ("预期与进展",
     "问题已由真机执行剖面量化（下页）；方案分层设计完成，对照实验进行中。",
     None, True),
], y0=Inches(1.75), row_h=Inches(1.18), gap=Inches(0.16), width=ROWS_W)
side_figure(s, "figure-rp2-idea.png")
footer(s, 9)
set_notes(s, "关键在于双工负载的输入形状是固定的：每个周期进多少输入由帧率和周期长度决定，一成不变，执行形状只有有限几种。但通用引擎按「这一步有没有新输入」来选执行路径，结构性看不见这种稳定性——它的优化机制在我们的负载上系统性落空，这是真机剖面量化过的。思路是为每种形状预编译静态执行计划，运行时直接重放，失配就回退，不影响正确性。")

# ================= 第 10 页：研究点 2 · 实验验证 =================
s = new_slide()
content_header(s, "研究点 2｜实验验证", "执行剖面实测：收益从哪里来", kicker_color=AMBER)
FIG_EV = os.path.join(os.path.dirname(os.path.abspath(__file__)), "figure-rp2-evidence.png")
if os.path.exists(FIG_EV):
    from PIL import Image as _PILImage
    _w, _h = _PILImage.open(FIG_EV).size
    fw = CONTENT_W
    fh = int(fw * _h / _w)
    s.shapes.add_picture(FIG_EV, MARGIN, Inches(1.78), fw, fh)
block_rows(s, [
    ("实测发现",
     "每周期的关键路径步恰好未被引擎图化机制覆盖，耗时约为已图化步的 3 倍；周期几乎被执行占满，后台结果没有可用的写回空隙。",
     None, False),
    ("预期与边界",
     "特化后推导可腾出约四分之一周期，数千 token 结果一至两个周期写回；收益为推导预期，对照实验进行中（消费级 GPU 初步实验）。",
     None, True),
], y0=Inches(5.62), row_h=Inches(0.62), gap=Inches(0.12), label_w=Inches(1.6), body_size=12.5)
footer(s, 10)
set_notes(s, "这一页讲研究点二的收益从哪里来。上面一条是实测的周期形状：橙色描边的是每周期八个关键步，它们恰好落在引擎图化机制的覆盖之外，每个约三倍于已图化的常规步——整个周期被执行占满，后台结果想写回，没有空隙。下面一条是特化后的推导形状：关键步纳入整步重放后收缩，周期尾部腾出约四分之一的绿色预算，数千 token 的后台结果一两个周期就能写回。对 JiuwenSwarm 这样的系统，这决定了后台干完的活能不能及时说出来。边界照旧：推导预期，对照实验进行中。")

# ================= 第 11 页：研究点 3 =================
s = new_slide()
content_header(s, "研究点 3｜立项设想", "端到端交互链路的联合调度", kicker_color=GRAY)
block_rows(s, [
    ("研究问题",
     "双工把交互变成持续过程：deadline 逐周期重复，网络、推理、后台、播放任何一环的抖动都会周期性撞上时限；各环节只能按最坏情况各自留余量，余量层层叠加，最终由并发容量买单。",
     None, False),
    ("关键洞察",
     "链路上唯一的松弛在弹性负载：后台调用与结果写回天然容忍延迟。但其时机横跨传输、引擎、播放三个子系统，任何单一环节都看不到全局——协调只能发生在一层联合调度中。",
     None, False),
    ("技术路线",
     "以统一交互时间轴为共享状态：推理感知网络抖动，写回给 deadline 让路，播放同时吸收网络与模型两个方差源。",
     None, False),
    ("初步观察",
     "已有实测实例：单环节的优化会被相邻环节的资源竞争抵消——孤立优化不可行并非假设；将复用研究点 1、2 的实验设施验证。",
     None, True),
], y0=Inches(1.75), row_h=Inches(1.18), gap=Inches(0.16), label_w=Inches(1.6), width=ROWS_W)
side_figure(s, "figure-rp3-idea.png")
footer(s, 11)
set_notes(s, "研究点三一句话：一次双工交互要过网络、前台推理、后台智能体、播放四个环节，今天各管各的、层层留余量，端到端时限没人统筹。我们的洞察是负载有刚有弹——前台周期和播放是刚性 deadline，后台调用和写回天然容忍延迟，刚弹分明才给联合调度留出空间。做法是让四个环节共享同一条交互时间轴：推理感知网络、写回给 deadline 让路、播放吸收两头的方差。这不是空想，实验里已经出现过单环节优化被相邻环节抵消的实例；立项后复用前两个研究点的设施验证。")

# ================= 第 12 页：交付物 =================
s = new_slide()
content_header(s, "预期成果", "交付物：可直接部署的双工交互推理引擎与接入层")
block_rows(s, [
    ("推理引擎",
     "一个面向双工交互负载的推理服务引擎：在 vLLM 上以增量补丁实现状态虚拟化、静态执行特化与时间轴调度，各机制可独立开关、可消融，随开源社区版本演进。",
     None, True),
    ("智能体接入层",
     "统一交互时间轴协议 + 会话接入与后台结果写回接口：JiuwenSwarm 等 agentic 系统无需改动内核即可接入，获得实时双工语音交互能力。",
     None, False),
    ("评测与观测套件",
     "双工负载的评测协议与全链路白盒观测：负载生成、公平对照方法、逐周期时间线，实验结论可复现、可审计。",
     None, False),
    ("研究产出",
     "系统方向论文与技术报告；按研究点分阶段交付——状态虚拟化、执行特化、联合调度，每项均以真机实验验证收尾。",
     None, False),
], y0=Inches(1.85), row_h=Inches(1.13), gap=Inches(0.18), label_w=Inches(1.75))
footer(s, 12)
set_notes(s, "最后落到交付物：不是一个概念，是四样具体的东西。核心是推理引擎——在 vLLM 上以增量补丁实现三个研究点的机制，每个机制可独立开关、可消融；配一个智能体接入层，JiuwenSwarm 这样的系统不改内核就能接入获得双工语音能力；评测与观测套件保证所有结论可复现；研究产出按研究点分阶段推进，每一阶段都以真机实验收尾。")

out = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                   "omni-anything-proposal.pptx")
prs.save(out)
print("saved:", out)
